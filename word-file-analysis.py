from docx import Document

# 워드 파일 열기
doc = Document("11-24-0633-15-00bn-mar-may-tgbn-teleconference-agenda.docx")

# print(doc.paragraphs)
submission_list = []
for table in doc.tables[2:3]:
    for row in table.rows[3:15]:
        items = []
        for cell in row.cells:
            items.append(cell.text)
        submission_list.append(items)
print(submission_list)

# 모든 단락 읽기
# for para in doc.paragraphs:
#     print(para.text)

# # 테이블 데이터 읽기
# for table in doc.tables:
#     for row in table.rows:
#         for cell in row.cells:
#             print(cell.text)
